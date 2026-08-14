import csv
import json
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.config import STORAGE_DIR
from services.journey_map_service import list_journeys


_ARTIFACT_PATHS = {
    "user-stories": STORAGE_DIR / "user_stories.json",
    "test-cases": STORAGE_DIR / "test_cases.json",
    "test-scripts": STORAGE_DIR / "test_scripts.json",
}


def _read_items(artifact_type: str) -> list[dict]:
    path = _ARTIFACT_PATHS.get(artifact_type)
    if not path or not path.exists():
        return []
    raw = path.read_text(encoding="utf-8")
    items = json.loads(raw) if raw.strip() else []
    return [item for item in items if isinstance(item, dict)] if isinstance(items, list) else []


def _selected_pipeline(journey_ids: list[str] | None = None) -> dict:
    journeys = [item for item in list_journeys() if isinstance(item, dict)]
    active_ids = {str(item.get("journey_id") or "") for item in journeys}
    requested_ids = {str(value) for value in (journey_ids or []) if value}
    selected_ids = (requested_ids & active_ids) if requested_ids else active_ids
    selected_journeys = [item for item in journeys if str(item.get("journey_id") or "") in selected_ids]

    stories = [
        item for item in _read_items("user-stories")
        if str(item.get("journey_id") or "") in selected_ids
    ]
    story_ids = {str(item.get("story_id") or "") for item in stories if item.get("story_id")}

    test_cases = [
        item for item in _read_items("test-cases")
        if str(item.get("journey_id") or "") in selected_ids
        and (not item.get("user_story_id") or str(item.get("user_story_id")) in story_ids)
    ]
    test_case_ids = {str(item.get("test_case_id") or "") for item in test_cases if item.get("test_case_id")}

    test_scripts = [
        item for item in _read_items("test-scripts")
        if str(item.get("test_case_id") or "") in test_case_ids
    ]

    return {
        "journeys": selected_journeys,
        "user-stories": stories,
        "test-cases": test_cases,
        "test-scripts": test_scripts,
    }


def _journey_preview(items: list[dict]) -> list[dict]:
    return [{
        "id": item.get("journey_id", ""),
        "journey_id": item.get("journey_id", ""),
        "title": item.get("journey_title") or item.get("starting_point") or "Untitled journey",
        "source_url": item.get("starting_url") or item.get("app_url") or "",
        "outcome": item.get("outcome") or "Captured",
        "captured_at": (item.get("exploration_metadata") or {}).get("exploration_timestamp") or item.get("created_at") or "",
        "objective": (item.get("test_hints") or {}).get("journey_objective") or (item.get("exploration_metadata") or {}).get("task_input") or "",
        "step_count": len(item.get("steps") or []),
        "event_count": len(item.get("events") or item.get("event_log") or []),
    } for item in items]


def _story_preview(items: list[dict]) -> list[dict]:
    return [{
        "id": item.get("story_id", ""),
        "journey_id": item.get("journey_id", ""),
        "title": item.get("summary") or "Untitled user story",
        "module": item.get("module") or item.get("epic") or "",
        "description": item.get("description") or "",
        "acceptance_criteria": item.get("acceptance_criteria") or [],
        "status": item.get("jira_sync_status") or "Local",
    } for item in items]


def _test_case_preview(items: list[dict]) -> list[dict]:
    return [{
        "id": item.get("test_case_id", ""),
        "journey_id": item.get("journey_id", ""),
        "title": item.get("title") or "Untitled test case",
        "description": item.get("description") or "",
        "type": item.get("test_case_type") or item.get("scenario_type") or "Functional",
        "steps": item.get("steps") or [],
    } for item in items]


def _test_script_preview(items: list[dict]) -> list[dict]:
    return [{
        "id": item.get("script_id", ""),
        "test_case_id": item.get("test_case_id", ""),
        "feature_filename": item.get("feature_filename") or "",
        "javascript_filename": item.get("javascript_filename") or "",
    } for item in items]


def _build_preview(artifact_type: str, items: list[dict]) -> list[dict]:
    if artifact_type == "journeys":
        return _journey_preview(items)
    if artifact_type == "user-stories":
        return _story_preview(items)
    if artifact_type == "test-cases":
        return _test_case_preview(items)
    if artifact_type == "test-scripts":
        return _test_script_preview(items)
    return []


def export_artifact(artifact_type: str, preview: bool = False, journey_ids: list[str] | None = None) -> dict:
    pipeline = _selected_pipeline(journey_ids)
    selected_ids = [str(item.get("journey_id") or "") for item in pipeline["journeys"]]

    if artifact_type == "complete-package":
        return {
            "artifact_type": artifact_type,
            "journey_ids": selected_ids,
            "counts": {key: len(value) for key, value in pipeline.items()},
            "journeys": pipeline["journeys"],
            "user_stories": pipeline["user-stories"],
            "test_cases": pipeline["test-cases"],
            "test_scripts": pipeline["test-scripts"],
        }

    items = pipeline.get(artifact_type, [])
    output_items = _build_preview(artifact_type, items) if preview else items
    return {
        "artifact_type": artifact_type,
        "journey_ids": selected_ids,
        "items": output_items,
        "count": len(output_items),
        "preview": preview,
    }


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _write_csv(path: Path, rows: list[dict], columns: list[tuple[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=[heading for heading, _ in columns])
        writer.writeheader()
        for item in rows:
            writer.writerow({heading: _text(item.get(key)) for heading, key in columns})


def _safe_archive_name(value: str, fallback: str) -> str:
    name = Path(str(value or fallback)).name
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name) or fallback


def _excel_value(value) -> str | int | float | bool:
    if isinstance(value, (int, float, bool)):
        return value
    text = _text(value)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)
    return text[:32767]


def _add_sheet(workbook: Workbook, title: str, rows: list[dict], columns: list[tuple[str, str]]) -> None:
    sheet = workbook.create_sheet(title=title)
    headings = [heading for heading, _ in columns]
    sheet.append(headings)
    for item in rows:
        sheet.append([_excel_value(item.get(key)) for _, key in columns])
    header_fill = PatternFill("solid", fgColor="FF6600")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, heading in enumerate(headings, start=1):
        sample_lengths = [len(str(sheet.cell(row=row, column=index).value or "")) for row in range(1, min(sheet.max_row, 30) + 1)]
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(max(len(heading) + 2, max(sample_lengths, default=10) + 2), 48)


STORY_COLUMNS = [
    ("Story ID", "story_id"), ("Journey ID", "journey_id"), ("User Story", "summary"),
    ("Description", "description"), ("Epic", "epic"), ("Module", "module"),
    ("Acceptance Criteria", "acceptance_criteria"), ("Jira Key", "jira_key"),
    ("Jira URL", "jira_url"), ("Jira Status", "jira_sync_status"), ("Generated At", "generated_at"),
]

CASE_COLUMNS = [
    ("Test Case ID", "test_case_id"), ("Journey ID", "journey_id"), ("User Story ID", "user_story_id"),
    ("Title", "title"), ("Description", "description"), ("Type", "test_case_type"),
    ("Priority", "priority"), ("Preconditions", "preconditions"), ("Steps", "steps"),
    ("Expected Results", "expected_results"), ("Generated At", "generated_at"),
]


def create_export_file(artifact_type: str, output_dir: Path, journey_ids: list[str] | None = None) -> tuple[Path, str, str]:
    pipeline = _selected_pipeline(journey_ids)
    output_dir.mkdir(parents=True, exist_ok=True)

    if artifact_type == "user-stories":
        filename = "user_stories.csv"
        path = output_dir / filename
        _write_csv(path, pipeline["user-stories"], STORY_COLUMNS)
        return path, filename, "text/csv; charset=utf-8"

    if artifact_type == "test-cases":
        filename = "test_cases.csv"
        path = output_dir / filename
        _write_csv(path, pipeline["test-cases"], CASE_COLUMNS)
        return path, filename, "text/csv; charset=utf-8"

    if artifact_type == "test-scripts":
        filename = "test_scripts_feature_and_js.zip"
        path = output_dir / filename
        with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
            for index, script in enumerate(pipeline["test-scripts"], start=1):
                script_id = str(script.get("script_id") or f"script_{index}")
                feature_name = _safe_archive_name(script.get("feature_filename"), f"{script_id}.feature")
                javascript_name = _safe_archive_name(script.get("javascript_filename"), f"{script_id}.js")
                archive.writestr(f"features/{feature_name}", _text(script.get("feature_file")))
                archive.writestr(f"javascript/{javascript_name}", _text(script.get("javascript_file")))
        return path, filename, "application/zip"

    if artifact_type == "complete-package":
        filename = "qa_complete_package.xlsx"
        path = output_dir / filename
        workbook = Workbook()
        workbook.remove(workbook.active)
        journey_rows = [{
            **item,
            "objective": (item.get("test_hints") or {}).get("journey_objective") or (item.get("exploration_metadata") or {}).get("task_input") or "",
            "step_count": len(item.get("steps") or []),
            "event_count": len(item.get("events") or item.get("event_log") or []),
            "captured_at": (item.get("exploration_metadata") or {}).get("exploration_timestamp") or item.get("created_at") or "",
        } for item in pipeline["journeys"]]
        _add_sheet(workbook, "Journeys", journey_rows, [
            ("Journey ID", "journey_id"), ("Journey", "journey_title"), ("Starting Point", "starting_point"),
            ("Starting URL", "starting_url"), ("Objective", "objective"), ("Outcome", "outcome"),
            ("Steps", "step_count"), ("Events", "event_count"), ("Captured At", "captured_at"),
        ])
        _add_sheet(workbook, "User Stories", pipeline["user-stories"], STORY_COLUMNS)
        _add_sheet(workbook, "Test Cases", pipeline["test-cases"], CASE_COLUMNS)
        _add_sheet(workbook, "Test Scripts", pipeline["test-scripts"], [
            ("Script ID", "script_id"), ("Test Case ID", "test_case_id"),
            ("Feature Filename", "feature_filename"), ("JavaScript Filename", "javascript_filename"),
            ("Gherkin", "feature_file"), ("JavaScript", "javascript_file"),
            ("Source Evidence Count", "source_evidence_count"), ("Generated At", "generated_at"),
        ])
        workbook.save(path)
        return path, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    raise ValueError(f"Unsupported export type: {artifact_type}")
